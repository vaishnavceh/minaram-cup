"""
Export service generating professionally styled XLSX and CSV reports
for registrations, teams, fixtures, results, standings, and complete summaries.
"""

import csv
import logging
from pathlib import Path
from typing import List, Dict, Any, Optional

from app.database import Database
from app.standings import StandingsCalculator
from app.models import Registration, Team, Match

logger = logging.getLogger("minaram_cup.export_service")


class ExportService:
    def __init__(self, db: Database):
        self.db = db
        self.standings = StandingsCalculator(db)

    # ----------------------------------------------------
    # Excel Styling Helper
    # ----------------------------------------------------
    @staticmethod
    def _apply_excel_styling(sheet, header_bg: str = "1A242F", header_fg: str = "FFFFFF"):
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        header_font = Font(name="Segoe UI", size=11, bold=True, color=header_fg)
        header_fill = PatternFill("solid", fgColor=header_bg)
        thin_border = Border(
            left=Side(style='thin', color="E2E8F0"),
            right=Side(style='thin', color="E2E8F0"),
            top=Side(style='thin', color="E2E8F0"),
            bottom=Side(style='thin', color="E2E8F0")
        )

        sheet.views.sheetView[0].showGridLines = True
        sheet.freeze_panes = "A2"

        # Header row
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Body rows
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                cell.font = Font(name="Segoe UI", size=10)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        # Auto-adjust column widths
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ----------------------------------------------------
    # Export Registrations
    # ----------------------------------------------------
    def export_registrations(self, file_path: Path, status_filter: Optional[str] = None) -> None:
        path = Path(file_path)
        query = "SELECT id, participant_name, phone_number, team_name, registered_at, status, notes FROM registrations"
        params = []
        if status_filter and status_filter != "All":
            query += " WHERE status = ?"
            params.append(status_filter)
        query += " ORDER BY id ASC;"
        rows = self.db.execute_query(query, tuple(params))

        headers = ["Reg ID", "Participant Name", "Phone Number", "Team Name", "Registered At", "Status", "Notes"]

        if path.suffix.lower() == ".csv":
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                for r in rows:
                    writer.writerow([r["id"], r["participant_name"], f"'{r['phone_number']}", r["team_name"], r["registered_at"], r["status"], r["notes"]])
        else:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Registrations"
            ws.append(headers)
            for r in rows:
                ws.append([r["id"], r["participant_name"], str(r["phone_number"]), r["team_name"], r["registered_at"], r["status"], r["notes"]])
            self._apply_excel_styling(ws)
            wb.save(path)

    # ----------------------------------------------------
    # Export Teams
    # ----------------------------------------------------
    def export_teams(self, file_path: Path) -> None:
        path = Path(file_path)
        rows = self.db.execute_query("""
            SELECT t.id, t.name, t.captain_name, t.phone_number, t.status, g.name AS group_name, t.seed, t.notes
            FROM teams t
            LEFT JOIN groups g ON t.group_id = g.id
            ORDER BY t.seed DESC, t.name ASC;
        """)
        headers = ["Team ID", "Team Name", "Captain", "Contact Phone", "Status", "Group", "Seed", "Notes"]

        if path.suffix.lower() == ".csv":
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                for r in rows:
                    writer.writerow([r["id"], r["name"], r["captain_name"], f"'{r['phone_number']}", r["status"], r["group_name"] or "None", r["seed"], r["notes"]])
        else:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Teams"
            ws.append(headers)
            for r in rows:
                ws.append([r["id"], r["name"], r["captain_name"], str(r["phone_number"]), r["status"], r["group_name"] or "None", r["seed"], r["notes"]])
            self._apply_excel_styling(ws)
            wb.save(path)

    # ----------------------------------------------------
    # Export Fixtures & Results
    # ----------------------------------------------------
    def export_fixtures(self, file_path: Path) -> None:
        path = Path(file_path)
        rows = self.db.execute_query("""
            SELECT m.match_number, m.stage, g.name AS group_name, m.round_number,
                   ht.name AS home_team, at.name AS away_team,
                   m.scheduled_date, m.scheduled_time, m.venue,
                   m.home_score, m.away_score, m.home_penalty, m.away_penalty,
                   m.status, wt.name AS winner_name
            FROM matches m
            LEFT JOIN groups g ON m.group_id = g.id
            LEFT JOIN teams ht ON m.home_team_id = ht.id
            LEFT JOIN teams at ON m.away_team_id = at.id
            LEFT JOIN teams wt ON m.winner_team_id = wt.id
            ORDER BY m.match_number ASC;
        """)
        headers = ["Match #", "Stage", "Group", "Round", "Home Team", "Away Team", "Date", "Time", "Venue", "Score", "Status", "Winner"]

        def fmt_score(r):
            if r["home_score"] is None or r["away_score"] is None:
                return "vs"
            s = f"{r['home_score']} - {r['away_score']}"
            if r["home_penalty"] is not None and r["away_penalty"] is not None:
                s += f" ({r['home_penalty']}-{r['away_penalty']} pen)"
            return s

        if path.suffix.lower() == ".csv":
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                for r in rows:
                    writer.writerow([
                        r["match_number"], r["stage"], r["group_name"] or "-", r["round_number"],
                        r["home_team"] or "TBD", r["away_team"] or "TBD",
                        r["scheduled_date"], r["scheduled_time"], r["venue"],
                        fmt_score(r), r["status"], r["winner_name"] or "-"
                    ])
        else:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Fixtures"
            ws.append(headers)
            for r in rows:
                ws.append([
                    r["match_number"], r["stage"], r["group_name"] or "-", r["round_number"],
                    r["home_team"] or "TBD", r["away_team"] or "TBD",
                    r["scheduled_date"], r["scheduled_time"], r["venue"],
                    fmt_score(r), r["status"], r["winner_name"] or "-"
                ])
            self._apply_excel_styling(ws)
            wb.save(path)

    # ----------------------------------------------------
    # Export Standings
    # ----------------------------------------------------
    def export_standings(self, file_path: Path) -> None:
        path = Path(file_path)
        all_groups = self.standings.calculate_all_groups_standings()
        headers = ["Rank", "Group", "Team", "Played", "Won", "Drawn", "Lost", "GF", "GA", "GD", "Points", "Qualified"]

        if path.suffix.lower() == ".csv":
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                for g_name, s_rows in all_groups.items():
                    for s in s_rows:
                        writer.writerow([s.rank, g_name, s.team_name, s.played, s.won, s.drawn, s.lost, s.goals_for, s.goals_against, s.goal_difference, s.points, "Yes" if s.is_qualified else "No"])
        else:
            from openpyxl import Workbook
            wb = Workbook()
            first = True
            for g_name, s_rows in all_groups.items():
                if first:
                    ws = wb.active
                    ws.title = f"Group {g_name}" if not g_name.startswith("Group") else g_name
                    first = False
                else:
                    sheet_title = f"Group {g_name}" if not g_name.startswith("Group") else g_name
                    ws = wb.create_sheet(title=sheet_title[:31])

                g_headers = ["Rank", "Team", "Played", "Won", "Drawn", "Lost", "GF", "GA", "GD", "Points", "Qualified"]
                ws.append(g_headers)
                for s in s_rows:
                    ws.append([s.rank, s.team_name, s.played, s.won, s.drawn, s.lost, s.goals_for, s.goals_against, s.goal_difference, s.points, "QUALIFIED" if s.is_qualified else ""])
                self._apply_excel_styling(ws)

            if first:  # If no groups existed
                ws = wb.active
                ws.title = "Standings"
                ws.append(headers)

            wb.save(path)

    # ----------------------------------------------------
    # Export Full Tournament Multi-Sheet Workbook
    # ----------------------------------------------------
    def export_full_tournament_workbook(self, file_path: Path) -> None:
        """Create a master Excel report with Overview, Teams, Standings, and Fixtures."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        path = Path(file_path)
        wb = Workbook()

        # 1. Overview Sheet
        ws_overview = wb.active
        ws_overview.title = "Tournament Overview"
        settings = self.db.get_all_settings()

        ws_overview["A1"] = f"{settings.get('tournament_name', 'Minaram Cup')} - Official Tournament Report"
        ws_overview["A1"].font = Font(name="Segoe UI", size=16, bold=True, color="10B981")
        ws_overview["A2"] = f"Edition: {settings.get('edition', '2026')} | Sport: {settings.get('sport', 'Football')} | Format: {settings.get('format_type', 'Group Stage + Knockout')}"
        ws_overview["A2"].font = Font(name="Segoe UI", size=11, italic=True)

        kpis = [
            ("Total Registrations", self.db.execute_one("SELECT COUNT(*) AS c FROM registrations;")["c"]),
            ("Confirmed Teams", self.db.execute_one("SELECT COUNT(*) AS c FROM teams WHERE status = 'Confirmed';")["c"]),
            ("Total Matches Scheduled", self.db.execute_one("SELECT COUNT(*) AS c FROM matches;")["c"]),
            ("Completed Matches", self.db.execute_one("SELECT COUNT(*) AS c FROM matches WHERE status IN ('Completed', 'Walkover');")["c"]),
            ("Win Points", settings.get("win_points", "3")),
            ("Draw Points", settings.get("draw_points", "1")),
            ("Loss Points", settings.get("loss_points", "0")),
            ("Default Venue", settings.get("default_venue", "Minaram Stadium")),
        ]

        ws_overview.append([])
        ws_overview.append(["Metric", "Value"])
        for k, v in kpis:
            ws_overview.append([k, v])
        self._apply_excel_styling(ws_overview, header_bg="1E293B")

        # 2. Teams Sheet
        ws_teams = wb.create_sheet(title="Teams")
        ws_teams.append(["Team ID", "Team Name", "Captain", "Contact Phone", "Status", "Group", "Seed"])
        teams_rows = self.db.execute_query("""
            SELECT t.id, t.name, t.captain_name, t.phone_number, t.status, g.name AS group_name, t.seed
            FROM teams t LEFT JOIN groups g ON t.group_id = g.id
            ORDER BY t.seed DESC, t.name ASC;
        """)
        for r in teams_rows:
            ws_teams.append([r["id"], r["name"], r["captain_name"], str(r["phone_number"]), r["status"], r["group_name"] or "-", r["seed"]])
        self._apply_excel_styling(ws_teams)

        # 3. Standings Sheet
        ws_standings = wb.create_sheet(title="Points Table")
        ws_standings.append(["Group", "Rank", "Team", "Played", "Won", "Drawn", "Lost", "GF", "GA", "GD", "Points", "Status"])
        all_groups = self.standings.calculate_all_groups_standings()
        for g_name, s_rows in all_groups.items():
            for s in s_rows:
                ws_standings.append([g_name, s.rank, s.team_name, s.played, s.won, s.drawn, s.lost, s.goals_for, s.goals_against, s.goal_difference, s.points, "QUALIFIED" if s.is_qualified else ""])
        self._apply_excel_styling(ws_standings)

        # 4. Fixtures & Results Sheet
        ws_fix = wb.create_sheet(title="Fixtures & Results")
        ws_fix.append(["Match #", "Stage", "Group", "Home Team", "Away Team", "Date", "Time", "Score", "Status", "Winner"])
        m_rows = self.db.execute_query("""
            SELECT m.match_number, m.stage, g.name AS group_name,
                   ht.name AS home_team, at.name AS away_team,
                   m.scheduled_date, m.scheduled_time,
                   m.home_score, m.away_score, m.home_penalty, m.away_penalty,
                   m.status, wt.name AS winner_name
            FROM matches m
            LEFT JOIN groups g ON m.group_id = g.id
            LEFT JOIN teams ht ON m.home_team_id = ht.id
            LEFT JOIN teams at ON m.away_team_id = at.id
            LEFT JOIN teams wt ON m.winner_team_id = wt.id
            ORDER BY m.match_number ASC;
        """)
        for r in m_rows:
            sc = "vs"
            if r["home_score"] is not None and r["away_score"] is not None:
                sc = f"{r['home_score']} - {r['away_score']}"
                if r["home_penalty"] is not None:
                    sc += f" ({r['home_penalty']}-{r['away_penalty']}p)"
            ws_fix.append([
                r["match_number"], r["stage"], r["group_name"] or "-",
                r["home_team"] or "TBD", r["away_team"] or "TBD",
                r["scheduled_date"], r["scheduled_time"], sc,
                r["status"], r["winner_name"] or "-"
            ])
        self._apply_excel_styling(ws_fix)

        wb.save(path)
